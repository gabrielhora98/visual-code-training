from openpyxl import load_workbook
import os

caminho = os.path.join(os.path.dirname(__file__), "WMS%20-%20KASHE_PYTHON.xlsx")

planilha = load_workbook(caminho)
print(planilha.sheetnames)


def buscar_produto(aba, produto_procurado):
    resultados = []

    for linha in aba.iter_rows(values_only=True):
        localizacao = linha[0]
        produto = linha[1]
        quantidade = linha[2]

        if produto == produto_procurado:
            resultados.append((localizacao, produto, quantidade))

    return resultados

def buscar_produto_promocao(aba, produto_procurado):
    resultados = []

    for linha in aba.iter_rows(values_only=True):
        produto = linha[0]
        quantidade = linha[1]

        if produto == produto_procurado:
            resultados.append((produto, quantidade))

    return resultados

def ver_estoque_completo(aba, localizacao_procurada):
    resultados = []

    for linha in aba.iter_rows(values_only=True):
        localizacao = linha[0]
        produto = linha[1]
        quantidade = linha[2]

        if localizacao == localizacao_procurada and quantidade > 0:
            resultados.append((localizacao, produto, quantidade))

    return resultados

while True:
    print("=== Estoque Kashe ===")
    print("1 - Consultar estoque")
    print("2 - Adicionar produto")
    print("3 - Retirar produto")
    print("4 - Sair")

    opcao = input("Digite a opção desejada: ")

    if opcao == "1":
        while True:
            print("=== Consultar estoque ===")
            print("1 - Ruas")
            print("2 - Salão")
            print("3 - Promoção")
            print("4 - Ver estoque completo")
            print("5 - Voltar")

            opcao = input("Escolha o estoque: ")

            if opcao == "1":
                produto_procurado = input("Digite o produto: ").strip().upper()

                resultados = []

                for nome_aba in ["Rua 1", "Rua 2", "Rua 3"]:
                    aba = planilha[nome_aba]

                    resultado = buscar_produto(aba, produto_procurado)

                    resultados.extend(resultado)
                if resultados:
                    for localizacao, produto, quantidade in resultados:
                        print("Localização:", localizacao)
                        print("Produto:", produto)
                        print("Quantidade:", quantidade)
                        print("----------------")
                else:
                    print("Produto não encontrado!")

            elif opcao == "2":
                produto_procurado = input("Digite o produto: ").strip().upper()
            
                aba = planilha["SALÃO"]

                resultados = buscar_produto(aba, produto_procurado)

                if resultados:
                    for localizacao, produto, quantidade in resultados:
                        print("Localização:", localizacao)
                        print("Produto:", produto)
                        print("Quantidade:", quantidade)
                        print("----------------")
                else:
                    print("Produto não encontrado!")


            elif opcao == "3":
                produto_procurado = input("Digite o produto que deseja procurar: ").strip().upper()

                aba = planilha["PROMOÇÃO"]

                resultados = buscar_produto_promocao(aba, produto_procurado)

                if resultados:
                    for produto, quantidade in resultados:
                        print("Produto:", produto)
                        print("Quantidade:", quantidade)
                        print("----------------")
                else:
                    print("Produto não encontrado!")
            elif opcao == "4":
                localizacao_procurada = input("Digite a localização que deseja consultar: ").strip()

                resultados = []

                for nome_aba in ["Rua 1", "Rua 2", "Rua 3"]:
                    aba = planilha[nome_aba]

                    resultado = ver_estoque_completo(aba, localizacao_procurada)

                    resultados.extend(resultado)

                if resultados:
                    for localizacao, produto, quantidade in resultados:
                        print("Localização:", localizacao)
                        print("Produto:", produto)
                        print("Quantidade:", quantidade)
                        print("----------------")
                else:
                    print("Localização não encontrada!")

            elif opcao == "5":
                break

            else:
                print("Opção inválida!")
    elif opcao == "2":
        while True:
            print("=== Adicionar produto ===")
            print("1 - Ruas")
            print("2 - Salão")
            print("3 - Promoção")
            print("4 - Voltar")

            opcao = input("Escolha o estoque: ")
            if opcao == "1":
                print("=== Ruas ===")
                print("1 - Rua 1")
                print("2 - Rua 2")
                print("3 - Rua 3")

                rua = input("Escolha a rua: ")
                if rua == "1":
                    aba = planilha["Rua 1"]

                elif rua == "2":
                    aba = planilha["Rua 2"]

                elif rua == "3":
                    aba = planilha["Rua 3"]

                else:
                    print("Rua inválida!")
                    continue

                produto = input("Digite o produto: ").strip().upper()
                quantidade = int(input("Digite a quantidade: "))
                localizacao = input("Digite a localização: ").strip()

                encontrado = False

                for numero_linha, linha in enumerate(aba.iter_rows(values_only=True), start=1):
                    local = linha[0]
                    produto_atual = linha[1]
                    quantidade_atual = linha[2]

                    if local == localizacao and produto_atual == produto:
                        nova_quantidade = quantidade_atual + quantidade

                        aba.cell(row=numero_linha, column=3).value = nova_quantidade

                        encontrado = True
                        break
                if not encontrado:
                    aba.append([localizacao, produto, quantidade])
                planilha.save(caminho)
                print("Produto adicionado com sucesso!")    

            elif opcao == "2":
                aba = planilha["SALÃO"]

                while True:
                    print("=== Salão ===")
                    print("1 - Arara 1")
                    print("2 - Arara 2")
                    print("3 - Arara 3")
                    print("4 - Mesa")
                    print("5 - Manequim")
                    print("6 - Voltar")

                    local_escolhido = input("Escolha o local: ")
                    if local_escolhido == "1":
                        local = "ARARA 1"

                    elif local_escolhido == "2":
                        local = "ARARA 2"

                    elif local_escolhido == "3":
                        local = "ARARA 3"

                    elif local_escolhido == "4":
                        local = "MESA"

                    elif local_escolhido == "5":
                        local = "MANEQUIM"

                    elif local_escolhido == "6":
                        break

                    else:
                        print("Opção inválida!")
                        continue
                    produto = input("Digite o produto: ").strip().upper()
                    quantidade = int(input("Digite a quantidade: "))

                    encontrado = False

                    for numero_linha, linha in enumerate(aba.iter_rows(values_only=True), start=1):
                        local_atual = linha[0]
                        produto_atual = linha[1]
                        quantidade_atual = linha[2]

                        if local_atual == local and produto_atual == produto:
                            nova_quantidade = quantidade_atual + quantidade

                            aba.cell(row=numero_linha, column=3).value = nova_quantidade

                            encontrado = True
                            break

                    if not encontrado:
                        aba.append([local, produto, quantidade])
                    planilha.save(caminho)
                    print("Produto adicionado ao salão com sucesso!")


            elif opcao == "3":
                aba = planilha["PROMOÇÃO"]

                produto = input("Digite o produto: ").strip().upper()
                quantidade = int(input("Digite a quantidade: "))  
                encontrado = False

                for numero_linha, linha in enumerate(aba.iter_rows(values_only=True), start=1):
                    produto_atual = linha[0]
                    quantidade_atual = linha[1]

                    if produto_atual == produto:
                        nova_quantidade = quantidade_atual + quantidade

                        aba.cell(row=numero_linha, column=2).value = nova_quantidade

                        encontrado = True
                        break
                if not encontrado:
                    aba.append([produto, quantidade])    
                planilha.save(caminho)
                print("Produto adicionado à promoção com sucesso!")

            elif opcao == "4":
                break

            else:
                print("Opção inválida!")

    elif opcao == "3":
        while True:
            print("=== Retirar produto ===")
            print("1 - Ruas")
            print("2 - Salão")
            print("3 - Promoção")
            print("4 - Voltar")

            opcao = input("Escolha o estoque: ")

            if opcao == "1":
                print("=== Ruas ===")
                print("1 - Rua 1")
                print("2 - Rua 2")
                print("3 - Rua 3")

                rua = input("Escolha a rua: ")
                if rua == "1":
                    aba = planilha["Rua 1"]

                elif rua == "2":
                    aba = planilha["Rua 2"]

                elif rua == "3":
                    aba = planilha["Rua 3"]

                else:
                    print("Rua inválida!")
                    continue
                produto = input("Digite o produto que deseja retirar: ").strip().upper()
                quantidade_retirar = int(input("Digite a quantidade que deseja retirar: "))
                localizacao = input("Digite a localização: ").strip()

                encontrado = False

                for numero_linha, linha in enumerate(aba.iter_rows(values_only=True), start=1):
                    local_atual = linha[0]
                    produto_atual = linha[1]
                    quantidade_atual = linha[2]

                    if local_atual == localizacao and produto_atual == produto:
                        encontrado = True

                        if quantidade_retirar <= quantidade_atual:
                            nova_quantidade = quantidade_atual - quantidade_retirar

                            aba.cell(row=numero_linha, column=3).value = nova_quantidade

                            print("Produto retirado com sucesso!")
                            print("Produto:", produto_atual)
                            print("Quantidade restante:", nova_quantidade)

                        else:
                            print("Quantidade insuficiente!")
                            print("Quantidade disponível:", quantidade_atual)
                        break
                if not encontrado:
                    print("Produto não encontrado nessa localização!")    
                planilha.save(caminho)

            elif opcao == "2":
                aba = planilha["SALÃO"]

                print("=== Salão ===")
                print("1 - Arara 1")
                print("2 - Arara 2")
                print("3 - Arara 3")
                print("4 - Mesa")
                print("5 - Manequim")

                local_escolhido = input("Escolha o local: ")
                if local_escolhido == "1":
                    local = "ARARA 1"

                elif local_escolhido == "2":
                    local = "ARARA 2"

                elif local_escolhido == "3":
                    local = "ARARA 3"

                elif local_escolhido == "4":
                    local = "MESA"

                elif local_escolhido == "5":
                    local = "MANEQUIM"

                else:
                    print("Opção inválida!")
                    continue
                produto = input("Digite o produto que deseja retirar: ").strip().upper()
                quantidade_retirar = int(input("Digite a quantidade que deseja retirar: "))
                encontrado = False

                for numero_linha, linha in enumerate(aba.iter_rows(values_only=True), start=1):
                    local_atual = linha[0]
                    produto_atual = linha[1]
                    quantidade_atual = linha[2]

                    if local_atual == local and produto_atual == produto:
                        encontrado = True

                        if quantidade_retirar <= quantidade_atual:
                            nova_quantidade = quantidade_atual - quantidade_retirar

                            aba.cell(row=numero_linha, column=3).value = nova_quantidade

                            print("Produto retirado com sucesso!")
                            print("Produto:", produto_atual)
                            print("Quantidade restante:", nova_quantidade)

                        else:
                            print("Quantidade insuficiente!")
                            print("Quantidade disponível:", quantidade_atual)
                        break
                if not encontrado:
                    print("Produto não encontrado nesse local!")
                planilha.save(caminho)


            elif opcao == "3":
                aba = planilha["PROMOÇÃO"]

                produto = input("Digite o produto que deseja retirar: ").strip().upper()
                quantidade_retirar = int(input("Digite a quantidade que deseja retirar: "))

                encontrado = False

                for numero_linha, linha in enumerate(aba.iter_rows(values_only=True), start=1):
                    produto_atual = linha[0]
                    quantidade_atual = linha[1]

                    if produto_atual == produto:
                        encontrado = True

                        if quantidade_retirar <= quantidade_atual:
                            nova_quantidade = quantidade_atual - quantidade_retirar

                            aba.cell(row=numero_linha, column=2).value = nova_quantidade

                            print("Produto retirado com sucesso!")
                            print("Produto:", produto_atual)
                            print("Quantidade restante:", nova_quantidade)

                        else:
                            print("Quantidade insuficiente!")
                            print("Quantidade disponível:", quantidade_atual)

                        break

                if not encontrado:
                    print("Produto não encontrado!")

                planilha.save(caminho)

            elif opcao == "4":
                break

            else:
                print("Opção inválida!")

    elif opcao == "4":
        print("Encerrando programa...")
        break

    else:
        print("Opção inválida!")
